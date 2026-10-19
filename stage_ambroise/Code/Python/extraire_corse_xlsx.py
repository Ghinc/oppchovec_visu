"""
Extrait les lignes de la Corse (départements 2A et 2B) d'une base INSEE communale au format xlsx (zippé ou non) et les écrit en CSV (séparateur « ; »).

Les bases du recensement 2023 (base_cc_*_2023_xlsx.zip, TD_*_2023_xlsx.zip) sont des xlsx de 30 à 80 Mo pour 35 000 communes : on les lit en flux (openpyxl, mode lecture seule),
sans les charger entièrement, et on ne garde que la Corse.

Nouveau format des bases 2023 : la première ligne contient des libellés en français (« Population (princ) », « Pop 0-14 ans (princ) »…) et non les codes habituels (P23_POP…).
La feuille « Métadonnées » donne la correspondance avec le « Code variable ancien format » (PNN_POP : NN = les deux derniers chiffres du millésime). Quand cette table existe,
les colonnes du CSV sont renommées avec les anciens codes, millésime 2023 (P23_POP) : le reste du code de calcul garde les mêmes noms de colonnes que pour 2016 et 2021.

Usage (depuis stage_ambroise/Code/Python/) :
    python extraire_corse_xlsx.py "<fichier.xlsx ou .zip>" "<sortie.csv>" [feuille] [millésime à 2 chiffres, défaut 23]
La feuille par défaut est la première dont le nom commence par COM (« COM_2023 »), sinon la première.
"""

import csv
import io
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

CLES = ("CODGEO", "GEO", "COM", "CODE", "CODGEO_2026", "CODGEO_2025", "Code géographique")


def correspondance(wb):
    """libellé -> ancien code, d'après la feuille « Métadonnées » ; {} si elle n'a pas ce format."""
    if "Métadonnées" not in wb.sheetnames:
        return {}
    table, dedans = {}, False
    for ligne in wb["Métadonnées"].iter_rows(values_only=True):
        if not dedans:
            dedans = bool(ligne) and str(ligne[0]).strip() == "Variable" and len(ligne) > 3 and "ancien format" in str(ligne[3])
            continue
        if ligne and ligne[0] and len(ligne) > 3 and ligne[3]:
            table[str(ligne[0]).strip()] = str(ligne[3]).strip()
    return table


def extraire(chemin, sortie, feuille=None, millesime="23"):
    chemin = Path(chemin)
    if chemin.suffix.lower() == ".zip":
        z = zipfile.ZipFile(chemin)
        nom = next(n for n in z.namelist() if n.lower().endswith(".xlsx"))
        source = io.BytesIO(z.read(nom))
    else:
        source = chemin
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb[feuille] if feuille else wb[next((n for n in wb.sheetnames if n.upper().startswith("COM")), wb.sheetnames[0])]
    table = correspondance(wb)
    entete, lignes, vues = None, [], 0
    for ligne in ws.iter_rows(values_only=True):
        vues += 1
        if entete is None:
            if ligne and ligne[0] is not None and str(ligne[0]).strip() in CLES:
                entete = [str(c).strip() if c is not None else "" for c in ligne]
            continue
        code = str(ligne[0]) if ligne and ligne[0] is not None else ""
        if code.startswith(("2A", "2B")):
            lignes.append(ligne)
    if entete is None:
        sys.exit(f"{chemin} : en-tête introuvable (feuilles : {wb.sheetnames})")
    renommees = 0
    if table:
        nouvel = [entete[0].replace("Code géographique", "CODGEO"), entete[1].replace("Libellé géographique", "LIBGEO")]
        for lib in entete[2:]:
            ancien = table.get(lib) or table.get(" ".join(lib.split()))
            nouvel.append(ancien.replace("NN", millesime, 1).replace("XX", millesime, 1) if ancien else lib)
            renommees += bool(ancien)
        entete = nouvel
    with open(sortie, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(entete)
        w.writerows(lignes)
    print(f"{chemin.name} : feuille {ws.title}, {vues} lignes lues, {len(lignes)} lignes de Corse, {len(entete)} colonnes dont {renommees} renommées -> {sortie}")


if __name__ == "__main__":
    if len(sys.argv) < 3 or len(sys.argv) > 5:
        sys.exit(__doc__)
    extraire(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None, sys.argv[4] if len(sys.argv) > 4 else "23")
